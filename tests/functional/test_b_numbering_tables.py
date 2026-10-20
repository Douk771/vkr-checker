# ruff: noqa: E501
"""SPEC 4.A: NUM-01..03, TBL-01..11 (без вёрстки: LibreOffice в тестовом окружении отсутствует)."""

import io

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from tests.functional.conftest import of_rule
from tests.functional.thesis import Thesis, p, table

BASE_TABLE = [
    ["Показатель", "Группа 1", "Группа 2"],
    ["Мотивация общая", "34,2", "38,5"],
    ["Адаптация общая", "41,0", "44,3"],
]


def viol(rep, rule):
    return [f for f in of_rule(rep, rule) if f["level"] == "violation"]


def susp(rep, rule):
    return [f for f in of_rule(rep, rule) if f["level"] == "suspicion"]


# ---- NUM-01 ---------------------------------------------------------------------------
def test_num01_gap_in_table_numbers(api, clean):
    rep = api.report(clean.sub("Таблица 2 – Сравнение групп", "Таблица 3 – Сравнение групп"))
    assert viol(rep, "NUM-01")


def test_num01_duplicate_table_number(api, clean):
    rep = api.report(clean.sub("Таблица 2 – Сравнение групп", "Таблица 1 – Сравнение групп"))
    assert viol(rep, "NUM-01")


def test_num01_wrong_order(api, clean):
    t = clean.sub("Таблица 1 – Описательные", "Таблица 2 – Описательные").sub(
        "Таблица 2 – Сравнение групп", "Таблица 1 – Сравнение групп"
    )
    assert viol(api.report(t), "NUM-01")


def test_num01_table_without_caption(api, clean):
    t = clean.no_caption("Таблица 2 – Сравнение групп")
    rep = api.report(t)
    assert any("подпис" in f["message"] for f in viol(rep, "NUM-01"))


def test_num01_figure_without_caption(api, clean):
    t = clean.no_caption("Рисунок 1 – Схема")
    rep = api.report(t)
    assert any("подпис" in f["message"] for f in viol(rep, "NUM-01"))


def test_num01_per_chapter_scheme_is_accepted(api, clean):
    t = (
        clean.sub("Таблица 1 – Описательные", "Таблица 2.1 – Описательные")
        .sub("Таблица 2 – Сравнение групп", "Таблица 2.2 – Сравнение групп")
        .sub("в табл. 1;", "в табл. 2.1;")
        .sub("в табл. 2,", "в табл. 2.2,")
        .sub("Рисунок 1 – Схема", "Рисунок 1.1 – Схема")
        .sub("см. рис. 1", "см. рис. 1.1")
    )
    assert not viol(api.report(t), "NUM-01")


def test_num01_mixed_schemes_is_violation(api, clean):
    t = clean.sub("Таблица 2 – Сравнение групп", "Таблица 2.2 – Сравнение групп").sub(
        "в табл. 2,", "в табл. 2.2,"
    )
    assert viol(api.report(t), "NUM-01")


def test_num01_appendix_own_numbering_allowed(api, clean):
    t = clean.insert_after(
        "Полные результаты диагностики.", table("Таблица А.1 – Сырые данные", BASE_TABLE)
    ).sub("полные данные даны в приложении А", "полные данные даны в приложении А (табл. А.1)")
    assert not viol(api.report(t), "NUM-01")


# ---- NUM-02 ---------------------------------------------------------------------------
def test_num02_table_without_reference(api, clean):
    t = clean.sub("приведены в табл. 1; ", "приведены ниже; ")
    fs = viol(api.report(t), "NUM-02")
    assert fs and any("1" in f["message"] for f in fs)


@pytest.mark.parametrize(
    "form", ["таблица 1", "(см. таблицу 1)", "табл. 1", "Таблицы 1", "в таблице 1"]
)
def test_num02_reference_forms_and_cases(api, clean, form):
    t = clean.sub("приведены в табл. 1;", f"приведены ({form});")
    assert not viol(api.report(t), "NUM-02")


@pytest.mark.parametrize(
    "form",
    [
        "рис. 1",
        "рисунке 1",
        "рисунку 1",
        "рисунки 1",
        "рисунок 1",
        "(см. рисунок 1)",
        "рисунка 1",
    ],
)
def test_num02_figure_reference_forms(api, clean, form):
    t = clean.sub("см. рис. 1", f"см. {form}")
    assert not viol(api.report(t), "NUM-02")


def test_num02_reference_to_nonexistent_number(api, clean):
    t = clean.sub("результаты в табл. 2,", "результаты в табл. 7,")
    assert viol(api.report(t), "NUM-02")


def test_num02_far_first_reference_is_suspicion(api, clean):
    """Первая ссылка далеко после таблицы: добавляем много текста между таблицей и единственной ссылкой."""
    from tests.functional.thesis import filler

    t = clean.sub("результаты в табл. 2,", "результаты ниже,")
    t = t.insert_after(
        "Гипотеза о том, что учебная",
        filler(60, "результатов"),
        p("Как видно из табл. 2, различия есть."),
    )
    rep = api.report(t)
    assert not viol(rep, "NUM-02")
    assert susp(rep, "NUM-02")


# ---- NUM-03 ---------------------------------------------------------------------------
def test_num03_count_mismatch_is_violation_with_both_values(api, clean):
    rep = api.report(clean.sub("1 рисунок и 2 таблицы", "3 рисунка и 2 таблицы"))
    fs = viol(rep, "NUM-03")
    assert fs and "3" in fs[0]["message"] and "1" in fs[0]["message"]


def test_num03_table_count_mismatch(api, clean):
    assert viol(api.report(clean.sub("1 рисунок и 2 таблицы", "1 рисунок и 5 таблиц")), "NUM-03")


def test_num03_phrase_missing_is_suspicion(api, clean):
    rep = api.report(
        clean.sub(
            "В работе содержится 1 рисунок и 2 таблицы, 1 приложение и 25 источников.", "Текст."
        )
    )
    assert susp(rep, "NUM-03") and not viol(rep, "NUM-03")


@pytest.mark.parametrize("verb", ["включает", "содержит"])
def test_num03_verb_variants(api, clean, verb):
    t = clean.sub("В работе содержится", f"Работа {verb}")
    assert not of_rule(api.report(t), "NUM-03")


def test_num03_sources_count_mismatch(api, clean):
    assert viol(api.report(clean.sub("25 источников", "40 источников")), "NUM-03")


def test_num03_appendix_tables_not_counted(api, clean):
    t = clean.insert_after(
        "Полные результаты диагностики.", table("Таблица А.1 – Сырые данные", BASE_TABLE)
    )
    assert not viol(api.report(t), "NUM-03")


# ---- TBL ------------------------------------------------------------------------------
def _patched(clean: Thesis, fn) -> bytes:
    d = Document(io.BytesIO(clean.bytes()))
    fn(d)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_tbl01_cell_indent_direct_is_violation(api, clean):
    def f(d):
        for t in d.tables:
            for row in t.rows:
                for c in row.cells:
                    c.paragraphs[0].paragraph_format.first_line_indent = Cm(1.25)

    assert viol(api.report(_patched(clean, f)), "TBL-01")


def test_tbl01_indent_inherited_from_normal_style_is_violation(api, clean):
    """SPEC: проверяется эффективное значение - красная строка из стиля «Обычный»."""

    def f(d):
        for t in d.tables:
            for row in t.rows:
                for c in row.cells:
                    pf = c.paragraphs[0].paragraph_format
                    pf._element.pPr.remove(pf._element.pPr.find(qn("w:ind")))

    assert viol(api.report(_patched(clean, f)), "TBL-01")


def test_tbl02_cell_spacing_is_violation(api, clean):
    def f(d):
        for t in d.tables:
            for row in t.rows:
                for c in row.cells:
                    c.paragraphs[0].paragraph_format.space_after = Pt(6)

    assert viol(api.report(_patched(clean, f)), "TBL-02")


def test_tbl02_line_spacing_one_and_half_is_violation(api, clean):
    from docx.enum.text import WD_LINE_SPACING

    def f(d):
        for t in d.tables:
            for row in t.rows:
                for c in row.cells:
                    c.paragraphs[
                        0
                    ].paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    assert viol(api.report(_patched(clean, f)), "TBL-02")


def test_tbl03_autofit_table_is_violation(api):
    assert viol(api.report(Thesis(autofit=True)), "TBL-03")


def test_tbl03_width_not_matching_text_width_is_violation(api):
    assert viol(api.report(Thesis(widths=[5, 3, 3])), "TBL-03")


def test_tbl03_width_within_tolerance_passes(api):
    assert not viol(api.report(Thesis(widths=[6.6, 5.0, 5.0])), "TBL-03")  # 16.6 см, допуск 0,2


def test_tbl05_header_left_aligned_is_violation(api, clean):
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    def f(d):
        for t in d.tables:
            for c in t.rows[0].cells:
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT

    assert viol(api.report(_patched(clean, f)), "TBL-05")


def test_tbl05_text_cells_justified_is_violation(api, clean):
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    def f(d):
        for t in d.tables:
            for row in t.rows[1:]:
                row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    assert viol(api.report(_patched(clean, f)), "TBL-05")


def test_tbl05_vertical_top_is_violation(api, clean):
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

    def f(d):
        for t in d.tables:
            for row in t.rows:
                for c in row.cells:
                    c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

    assert viol(api.report(_patched(clean, f)), "TBL-05")


def test_tbl07_exact_row_height_is_suspicion(api, clean):
    def f(d):
        for row in d.tables[0].rows:
            trPr = row._tr.get_or_add_trPr()
            h = OxmlElement("w:trHeight")
            h.set(qn("w:val"), str(int(5 * 567)))
            h.set(qn("w:hRule"), "exact")
            trPr.append(h)

    rep = api.report(_patched(clean, f))
    assert susp(rep, "TBL-07")


def test_tbl08_caption_without_keep_with_next(api, clean):
    t = clean.copy()
    for b in t.blocks:
        if b["k"] == "p" and b["t"].startswith("Таблица 1"):
            pass

    # подпись таблицы рендерится с keep_with_next; снимаем флаг в docx
    def f(d):
        for para in d.paragraphs:
            if para.text.startswith("Таблица 1 –"):
                para.paragraph_format.keep_with_next = False

    assert viol(api.report(_patched(clean, f)), "TBL-08")


def test_tbl09_table_with_missing_style(api, clean):
    def f(d):
        tblPr = d.tables[0]._tbl.tblPr
        st = tblPr.find(qn("w:tblStyle"))
        st.set(qn("w:val"), "ImportedFromElsewhere123")

    assert viol(api.report(_patched(clean, f)), "TBL-09")


def test_tbl10_long_table_in_main_text_is_suspicion(api, clean):
    rows = [["Показатель", "Группа 1", "Группа 2"]] + [
        [f"Параметр номер {i}", f"{i},1", f"{i},2"] for i in range(1, 60)
    ]
    t = (
        clean.insert_after(
            "Гипотеза о том, что учебная", table("Таблица 3 – Длинная таблица", rows)
        )
        .sub("1 рисунок и 2 таблицы", "1 рисунок и 3 таблицы")
        .insert_after("Гипотеза о том, что учебная", p("См. табл. 3 и приложение А."))
    )
    assert susp(api.report(t), "TBL-10")


def test_tbl11_stats_table_without_appendix_reference(api, clean):
    t = clean.sub("; полные данные даны в приложении А", "").sub(
        ", полные данные в приложении А", ""
    )
    assert susp(api.report(t), "TBL-11")
