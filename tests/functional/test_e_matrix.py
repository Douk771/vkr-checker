# ruff: noqa: E501
"""SPEC 4.D: MAT-01 (10 строк на размеченной работе, >=70% совпадений с эталоном, выгрузка .xlsx)."""

import io

import pytest
from openpyxl import load_workbook

from tests.functional.thesis import Thesis, matrix_blocks


@pytest.fixture(scope="module")
def mrep(api):
    return api.report(Thesis(matrix_blocks()))


def row_by_label(rep, label):
    return next(r for r in rep["matrix"]["rows"] if r["label"] == label)


def test_mat01_row_count_one_hypothesis_two_parts_three_positions_five_tasks(mrep):
    rows = mrep["matrix"]["rows"]
    assert len(rows) == 10 and mrep["matrix"]["stats"]["rows"] == 10
    kinds = [r["kind"] for r in rows]
    assert kinds.count("hypothesis_part") == 2
    assert kinds.count("defense_position") == 3
    assert kinds.count("task") == 5


# ожидаемые кандидаты (разметка составлена вместе с текстом работы, по SPEC 4.D)
EXPECT = {
    "Гипотеза, ч.1": dict(
        theory=("1.1", "1.2"),
        instrument=("тревожн", "самооценк"),
        calculation=("rs", "Спирмен"),
        conclusion=("вывод", "Вывод"),
    ),
    "Гипотеза, ч.2": dict(
        theory=("1.3",),
        instrument=("мотивац",),
        calculation=("U", "Манн"),
        conclusion=("вывод", "Вывод"),
    ),
    "Положение 1": dict(
        theory=("1.1", "1.2"),
        instrument=("тревожн", "самооценк"),
        calculation=("rs", "Спирмен"),
        conclusion=("вывод 4",),
    ),
    "Положение 2": dict(
        theory=("1.3",), instrument=("мотивац",), calculation=("U", "Манн"), conclusion=("вывод 5",)
    ),
    "Положение 3": dict(
        theory=("1.1",),
        instrument=("тревожн",),
        calculation=("Таблица 1",),
        conclusion=("вывод 2",),
    ),
    "Задача 1": dict(theory=("1.1", "1.2", "1.3"), conclusion=("вывод 1",)),
    "Задача 2": dict(
        theory=("1.1",),
        instrument=("тревожн",),
        calculation=("Таблица 1",),
        conclusion=("вывод 2",),
    ),
    "Задача 3": dict(
        theory=("1.2",),
        instrument=("самооценк",),
        calculation=("Таблица 2",),
        conclusion=("вывод 3",),
    ),
    "Задача 4": dict(
        theory=("1.1", "1.2"),
        instrument=("тревожн", "самооценк"),
        calculation=("rs", "Спирмен", "Таблица 3"),
        conclusion=("вывод 4",),
    ),
    "Задача 5": dict(
        theory=("1.3",),
        instrument=("мотивац",),
        calculation=("Таблица 4", "U", "Манн"),
        conclusion=("вывод 5",),
    ),
}


def test_mat01_candidates_match_reference_in_at_least_70_percent(mrep):
    total = hit = 0
    misses = []
    for label, exp in EXPECT.items():
        row = row_by_label(mrep, label)
        for cell, needles in exp.items():
            c = row["cells"][cell]
            blob = (c["text"] + " " + " ".join(c["candidates"])).lower()
            total += 1
            if c["status"] == "auto" and any(n.lower() in blob for n in needles):
                hit += 1
            else:
                misses.append((label, cell, c["text"][:60]))
    assert total >= 30
    assert hit / total >= 0.70, f"{hit}/{total}; промахи: {misses}"


def test_mat01_theoretical_task_has_no_instrument_and_calculation_gaps(mrep):
    r = row_by_label(mrep, "Задача 1")
    assert r["is_theoretical"]
    assert (
        r["cells"]["instrument"]["status"] == "n/a" and r["cells"]["calculation"]["status"] == "n/a"
    )


def test_mat01_empty_cell_is_marked_as_gap_and_counted(mrep):
    gaps = [
        (r["label"], k)
        for r in mrep["matrix"]["rows"]
        for k, c in r["cells"].items()
        if c["status"] == "gap"
    ]
    assert gaps and mrep["matrix"]["stats"]["gaps"] == len(gaps)
    n_a = [c for r in mrep["matrix"]["rows"] for c in r["cells"].values() if c["status"] == "n/a"]
    assert n_a  # n/a не считается разрывом


def test_mat01_cells_have_locations_for_auto_candidates(mrep):
    auto = [c for r in mrep["matrix"]["rows"] for c in r["cells"].values() if c["status"] == "auto"]
    assert auto and sum(1 for c in auto if c["locations"]) / len(auto) > 0.5


def test_mat01_matrix_present_but_empty_when_no_tasks(api, clean):
    t = clean.copy()
    t.blocks = [b for b in t.blocks if b["k"] != "auto"]
    rep = api.report(t)
    assert rep["matrix"]["stats"]["warnings"]  # понятное предупреждение, а не падение
    assert not [r for r in rep["matrix"]["rows"] if r["kind"] == "task"]


# ---- выгрузка .xlsx (UI-02) --------------------------------------------------------------
def export(client, rep, name="Моя работа.docx"):
    return client.post(
        "/api/export/matrix.xlsx",
        json={
            "file_name": name,
            "service_version": rep["service_version"],
            "checked_at": rep["checked_at"],
            "matrix": rep["matrix"],
        },
    )


def test_export_xlsx_is_editable_workbook_with_all_rows(client, mrep):
    r = export(client, mrep)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))  # открывается как обычная книга (редактируемая)
    ws = wb.active
    text = "\n".join(
        " | ".join(str(c.value) for c in row if c.value is not None) for row in ws.iter_rows()
    )
    for col in ("Теория", "Методика", "Расчёт", "Результат", "Вывод"):
        assert col.lower() in text.lower(), col
    assert "Задача 5" in text and "Гипотеза, ч.1" in text
    assert "проверить" in text.lower()  # «найдено автоматически, проверить»
    assert ws.max_row >= 11


def test_export_marks_gaps_visibly(client, mrep):
    ws = load_workbook(io.BytesIO(export(client, mrep).content)).active
    fills = {
        c.fill.fgColor.rgb for row in ws.iter_rows() for c in row if c.fill and c.fill.fill_type
    }
    texts = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert fills or "разрыв" in texts.lower() or "нет данных" in texts.lower()


def test_export_rejects_garbage_json(client):
    r = client.post("/api/export/matrix.xlsx", json={"matrix": "oops"})
    assert r.status_code == 422 and r.json()["error"]["code"] == "validation_error"


def test_export_rejects_non_json(client):
    r = client.post(
        "/api/export/matrix.xlsx", content=b"not json", headers={"content-type": "application/json"}
    )
    assert r.status_code in (400, 422)


def test_export_file_name_with_path_or_formula_is_safe(client, mrep):
    r = export(client, mrep, name='../../etc/passwd\n=HYPERLINK("http://x").docx')
    assert r.status_code in (200, 422)
    if r.status_code == 200:
        cd = r.headers["content-disposition"]
        assert "/" not in cd.split("filename", 1)[1].split(";")[0] and "\n" not in cd


def test_export_cell_text_starting_with_equals_is_not_a_formula(client, mrep):
    m = {**mrep["matrix"], "rows": [dict(mrep["matrix"]["rows"][0], statement="=1+1")]}
    r = client.post("/api/export/matrix.xlsx", json={"file_name": "x.docx", "matrix": m})
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.strip() == "=1+1":
                assert c.data_type != "f", "ячейка выгрузки стала формулой (инъекция формул)"
