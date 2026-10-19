# ruff: noqa: E501
"""Независимый от реализации сборщик синтетических ВКР для функциональных тестов.

Работа описывается списком блоков (dict); мутации правят блоки, а не файл, поэтому
пары «чистый / с ошибкой» отличаются ровно одной правкой. Текст целиком вымышленный.
Составлен по SPEC.md, а не по коду проверок.
"""

from __future__ import annotations

import copy
import io
import struct
import zlib
from collections.abc import Callable

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

NB = " "


def _png() -> bytes:
    def chunk(kind: bytes, data: bytes) -> bytes:
        body = kind + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body))

    raw = b"".join(b"\x00" + b"\x80" * 8 for _ in range(8))
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", 8, 8, 8, 0, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


# --- блоки ---------------------------------------------------------------------------
def h1(text, new_page=True):
    return {"k": "h1", "t": text, "np": new_page}


def h2(text):
    return {"k": "h2", "t": text}


def p(text, indent=True, align=None, keep=False, tag=None):
    return {"k": "p", "t": text, "indent": indent, "align": align, "keep": keep, "tag": tag}


def item(num, text):
    """Пункт нумерованного списка средствами Word (автонумерация, номера нет в тексте)."""
    return {"k": "auto", "t": text, "num": num}


def table(caption, rows, tag=None):
    return {"k": "table", "caption": caption, "rows": rows, "tag": tag}


def figure(caption):
    return {"k": "figure", "caption": caption}


def filler(n_par, topic="адаптационных процессов"):
    """Нейтральный текст без чисел, фамилий и ключевых слов проверок."""
    return {"k": "filler", "n": n_par, "topic": topic}


_ADJ = [
    "устойчивый",
    "гибкий",
    "широкий",
    "узкий",
    "глубокий",
    "заметный",
    "скрытый",
    "явный",
    "частный",
    "общий",
    "ранний",
    "поздний",
    "прочный",
    "хрупкий",
    "яркий",
    "сдержанный",
    "динамичный",
    "стабильный",
    "спорный",
    "типичный",
]
_NOUN = [
    "подход",
    "механизм",
    "признак",
    "фактор",
    "аспект",
    "процесс",
    "компонент",
    "критерий",
    "контекст",
    "уровень",
    "вариант",
    "принцип",
    "образец",
    "источник",
    "результат",
    "элемент",
    "параметр",
    "сценарий",
    "эпизод",
    "фрагмент",
]
_VERB = [
    "уточняется",
    "обсуждается",
    "описывается",
    "сопоставляется",
    "раскрывается",
    "фиксируется",
    "пересматривается",
    "иллюстрируется",
    "обобщается",
    "проверяется",
]
_COUNTER = [0]


def _uniq(topic: str) -> str:
    """Один абзац ~950 знаков с неповторяющимся содержимым (без чисел, фамилий, ключевых слов)."""
    out = []
    for _ in range(4):
        i = _COUNTER[0]
        _COUNTER[0] += 1
        out.append(
            f"Рассматривается {_ADJ[i % 20]} {_NOUN[(i // 20) % 20]} в контексте {topic}, "
            f"при этом {_ADJ[(i * 7 + 3) % 20]} {_NOUN[(i * 3 + i // 20) % 20]} {_VERB[(i // 7) % 10]} "
            f"с учётом условий образовательной среды и особенностей участников. "
            f"Такой порядок изложения сохраняет логику рассуждения и позволяет избежать повторов "
            f"при описании каждого последующего положения."
        )
    return " ".join(out)


def render_filler(n, topic, salt):
    return [_uniq(topic) for _ in range(n)]


# --- чистая работа -----------------------------------------------------------------
def clean_blocks(pages_main: int = 70) -> list[dict]:
    _COUNTER[0] = 0
    per = max((pages_main - 8) // 4, 1)  # страниц на подраздел
    paras = lambda pg: int(pg * 1800 / 950)  # noqa: E731
    b: list[dict] = []
    # титульный лист
    b += [
        p("МИНИСТЕРСТВО ОБРАЗОВАНИЯ (вымышленное)", False, "center"),
        p("Кафедра общей психологии", False, "center"),
        p("ВЫПУСКНАЯ КВАЛИФИКАЦИОННАЯ РАБОТА", False, "center"),
        p("Учебная мотивация и академическая адаптация студентов", False, "center"),
        p("Выполнил: студент группы М-21", False, "right"),
        p("Научный руководитель: доцент кафедры", False, "right"),
    ]
    toc = [
        "ВВЕДЕНИЕ",
        "ГЛАВА 1. Теоретические основы изучения учебной мотивации",
        "1.1. Понятие учебной мотивации",
        "1.2. Академическая адаптация студентов",
        "Выводы по главе 1",
        "ГЛАВА 2. Эмпирическое исследование мотивации и адаптации",
        "2.1. Организация и методы исследования",
        "2.2. Результаты исследования",
        "Выводы по главе 2",
        "ЗАКЛЮЧЕНИЕ",
        "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ",
        "ПРИЛОЖЕНИЕ А",
    ]
    b.append(h1("СОДЕРЖАНИЕ"))
    b += [p(f"{t} ........ {i + 3}", False, tag="toc") for i, t in enumerate(toc)]

    # введение
    b.append(h1("ВВЕДЕНИЕ"))
    b += [
        p(
            "Актуальность исследования определяется ростом интереса к адаптации студентов магистратуры."
        ),
        p(
            "Степень разработанности проблемы. Вопросы мотивации обучения подробно описаны в трудах "
            "С.Л. Рубинштейна и А.Н. Леонтьева."
        ),
        p(
            "Цель исследования: выявить связь учебной мотивации и академической адаптации студентов."
        ),
        p("Объект исследования: учебная мотивация студентов магистратуры."),
        p("Предмет исследования: связь учебной мотивации с уровнем академической адаптации."),
        p(
            "Гипотеза исследования: учебная мотивация студентов положительно связана с уровнем "
            "академической адаптации."
        ),
        p("Задачи исследования:", tag="tasks_head"),
        item(1, "Проанализировать литературу по проблеме учебной мотивации и адаптации."),
        item(2, "Подобрать методики диагностики учебной мотивации и академической адаптации."),
        item(3, "Выявить связь учебной мотивации с академической адаптацией студентов."),
        p(
            "Теоретико-методологические основания исследования составляют идеи деятельностного "
            "подхода А.Н. Леонтьева и теории мотивации С.Л. Рубинштейна."
        ),
        p(
            "Методы и методики исследования: опросник «Шкала учебной мотивации» (ШУМ), "
            "коэффициент корреляции Спирмена, U-критерий Манна–Уитни."
        ),
        p(
            f"Выборка исследования: n{NB}={NB}120 студентов магистратуры, из них 60% женщин и 40% мужчин."
        ),
        p(
            "Научная новизна исследования состоит в сопоставлении мотивации и адаптации на одной выборке."
        ),
        p(
            "Теоретическая значимость исследования заключается в уточнении представлений о связи "
            "мотивации и адаптации."
        ),
        p(
            "Практическая значимость исследования состоит в возможности использовать результаты "
            "в работе кураторов."
        ),
        p("Положения, выносимые на защиту:"),
        p("1. Мотивация студентов связана с их адаптацией.", False),
        p("2. Связь выражена сильнее у студентов первого года обучения.", False),
        p("3. Различия между группами по уровню мотивации статистически значимы.", False),
        p("Апробация результатов: доклад на научно-практической конференции."),
        p(
            "Структура работы. Работа состоит из введения, двух глав, заключения, списка источников "
            "и приложения. В работе содержится 1 рисунок и 2 таблицы, 1 приложение и 25 источников."
        ),
    ]
    # глава 1
    b.append(h1("ГЛАВА 1. Теоретические основы изучения учебной мотивации"))
    b.append(h2("1.1. Понятие учебной мотивации"))
    b += [
        p(
            "Шкала учебной мотивации разработана Н.Н. Ивановым и описана в литературе как инструмент "
            "диагностики. Подход С.Л. Рубинштейна и А.Н. Леонтьева рассматривает мотивацию как систему."
        ),
        filler(paras(per), "учебной мотивации"),
    ]
    b.append(h2("1.2. Академическая адаптация студентов"))
    b += [
        p(
            "Понятие адаптации раскрыто в работах П.П. Петрова, адаптировавшего шкалу для студентов. "
            "На рисунке ниже показана схема исследования, см. рис. 1."
        ),
        figure("Рисунок 1 – Схема исследования"),
        filler(paras(per), "академической адаптации"),
    ]
    b.append(h2("Выводы по главе 1"))
    b += [
        p(
            "Анализ литературы показал, что мотивация и адаптация рассматриваются как связанные явления."
        )
    ]
    # глава 2
    b.append(h1("ГЛАВА 2. Эмпирическое исследование мотивации и адаптации"))
    b.append(h2("2.1. Организация и методы исследования"))
    b += [
        p(
            f"В исследовании приняли участие 120 человек в возрасте от 18 до 25{NB}лет (M{NB}={NB}20,4; "
            f"SD{NB}={NB}1,8)."
        ),
        p(
            "Использованы методики: опросник «Шкала учебной мотивации» (ШУМ) Н.Н. Иванова, адаптация П.П. Петрова; "
            "для обработки данных применялись коэффициент корреляции Спирмена и U-критерий Манна–Уитни."
        ),
        filler(paras(per), "организации исследования"),
    ]
    b.append(h2("2.2. Результаты исследования"))
    b += [
        p("Описательные статистики приведены в табл. 1; полные данные даны в приложении А."),
        table(
            "Таблица 1 – Описательные статистики",
            [
                ["Показатель", "Группа 1", "Группа 2"],
                ["Среднее значение (M)", "12,4", "13,1"],
                ["Стандартное отклонение", "2,1", "2,4"],
            ],
        ),
        p(
            f"Выявлена положительная связь мотивации и адаптации (rs{NB}={NB}0,45; p{NB}<{NB}0,01), "
            f"различия между группами значимы (U{NB}={NB}120,5; p{NB}<{NB}0,05), результаты в табл. 2, полные данные в приложении А."
        ),
        table(
            "Таблица 2 – Сравнение групп",
            [
                ["Показатель", "Группа 1", "Группа 2"],
                ["Мотивация", "34,2", "38,5"],
                ["Адаптация", "41,0", "44,3"],
            ],
        ),
        p(
            "Гипотеза о том, что учебная мотивация студентов положительно связана с уровнем академической адаптации, подтверждена."
        ),
        filler(paras(per), "результатов исследования"),
    ]
    b.append(h2("Выводы по главе 2"))
    b += [
        p(
            "Данные эмпирического исследования показывают согласованность мотивации и адаптации у 120 студентов."
        )
    ]
    # заключение
    b.append(h1("ЗАКЛЮЧЕНИЕ"))
    b += [
        p("В заключении подведены итоги работы."),
        p(
            "1. Анализ литературы по проблеме учебной мотивации и адаптации показал общие основания их изучения.",
            False,
        ),
        p(
            "2. Для диагностики подобраны методики мотивации и адаптации, пригодные для студентов.",
            False,
        ),
        p(
            "3. Выявлена связь учебной мотивации с академической адаптацией: гипотеза о том, что учебная "
            "мотивация студентов положительно связана с уровнем академической адаптации, подтверждена "
            "на выборке из 120 студентов.",
            False,
        ),
    ]
    # литература
    b.append(h1("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ"))
    b += [
        p(
            f"{i}. Автор{chr(1040 + i % 20)} А.А. Название работы номер {i}. М.: Издательство, {2000 + i}. {100 + i} с.",
            False,
        )
        for i in range(1, 26)
    ]
    # приложение
    b.append(h1("ПРИЛОЖЕНИЕ А"))
    b += [p("Полные результаты диагностики.", False, "center")]
    return b


# --- рендер ----------------------------------------------------------------------------
def _align(name):
    return {
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
    }.get(name)


def _setup(doc):
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    n = doc.styles["Normal"]
    n.font.name = "Times New Roman"
    n.font.size = Pt(14)
    n.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    n.paragraph_format.first_line_indent = Cm(1.25)
    n.paragraph_format.space_before = n.paragraph_format.space_after = Pt(0)
    n.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name in ("Heading 1", "Heading 2"):
        st = doc.styles[name]
        st.font.name = "Times New Roman"
        st.font.size = Pt(14)
        st.font.bold = True
        st.font.color.rgb = None
        st.paragraph_format.first_line_indent = Cm(0)
        st.paragraph_format.keep_with_next = True


def _numbering(doc):
    """Список «1., 2., 3.» на уровне Word (numPr) без номера в тексте абзаца."""
    part = doc.part.numbering_part.element
    absn = OxmlElement("w:abstractNum")
    absn.set(qn("w:abstractNumId"), "90")
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    for tag, val in (
        ("w:start", "1"),
        ("w:numFmt", "decimal"),
        ("w:lvlText", "%1."),
        ("w:lvlJc", "left"),
    ):
        e = OxmlElement(tag)
        e.set(qn("w:val"), val)
        lvl.append(e)
    absn.append(lvl)
    part.insert(0, absn)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), "90")
    a = OxmlElement("w:abstractNumId")
    a.set(qn("w:val"), "90")
    num.append(a)
    part.append(num)


def _fixed_table(t, widths_cm):
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    for w in tblPr.findall(qn("w:tblW")):
        tblPr.remove(w)
    tw = OxmlElement("w:tblW")
    tw.set(qn("w:w"), str(int(sum(widths_cm) * 567)))
    tw.set(qn("w:type"), "dxa")
    tblPr.append(tw)
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_cm, strict=False):
        gc.set(qn("w:w"), str(int(w * 567)))
    for row in t.rows:
        for c, w in zip(row.cells, widths_cm, strict=False):
            c.width = Cm(w)


def _is_num(s):
    return s.replace(",", "").replace(".", "").replace("-", "").isdigit()


def render(blocks: list[dict], *, table_opts: dict | None = None) -> bytes:
    opts = table_opts or {}
    _COUNTER[0] = 0
    doc = Document()
    _setup(doc)
    _numbering(doc)
    for blk in blocks:
        k = blk["k"]
        if k in ("h1", "h2"):
            para = doc.add_heading(blk["t"], level=1 if k == "h1" else 2)
            if k == "h1":
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.page_break_before = blk["np"]
        elif k == "p":
            para = doc.add_paragraph(blk["t"])
            if not blk["indent"]:
                para.paragraph_format.first_line_indent = Cm(0)
            if blk["align"]:
                para.alignment = _align(blk["align"])
            if blk["keep"]:
                para.paragraph_format.keep_with_next = True
        elif k == "auto":
            para = doc.add_paragraph(blk["t"])
            ppr = para._p.get_or_add_pPr()
            numpr = OxmlElement("w:numPr")
            il = OxmlElement("w:ilvl")
            il.set(qn("w:val"), "0")
            ni = OxmlElement("w:numId")
            ni.set(qn("w:val"), "90")
            numpr.append(il)
            numpr.append(ni)
            ppr.append(numpr)
        elif k == "filler":
            for text in render_filler(blk["n"], blk["topic"], 0):
                doc.add_paragraph(text)
        elif k == "figure":
            holder = doc.add_paragraph()
            holder.alignment = WD_ALIGN_PARAGRAPH.CENTER
            holder.paragraph_format.first_line_indent = Cm(0)
            holder.add_run().add_picture(io.BytesIO(_png()), width=Cm(4))
            if blk["caption"]:
                cap = doc.add_paragraph(blk["caption"])
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.paragraph_format.first_line_indent = Cm(0)
        elif k == "table":
            if blk["caption"]:
                cap = doc.add_paragraph(blk["caption"])
                cap.paragraph_format.first_line_indent = Cm(0)
                cap.paragraph_format.keep_with_next = True
            rows = blk["rows"]
            t = doc.add_table(rows=len(rows), cols=len(rows[0]), style="Table Grid")
            ncols = len(rows[0])
            widths = opts.get("widths") or ([6.5] + [(16.5 - 6.5) / (ncols - 1)] * (ncols - 1))
            for r, row in enumerate(rows):
                for c, val in enumerate(row):
                    cell = t.cell(r, c)
                    cp = cell.paragraphs[0]
                    cp.add_run(val)
                    pf = cp.paragraph_format
                    pf.first_line_indent = Cm(opts.get("indent", 0))
                    pf.left_indent = Cm(0)
                    pf.space_after = Pt(0)
                    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
                    cp.alignment = (
                        WD_ALIGN_PARAGRAPH.CENTER
                        if (r == 0 or _is_num(val))
                        else WD_ALIGN_PARAGRAPH.LEFT
                    )
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if not opts.get("autofit"):
                _fixed_table(t, widths)
            doc.add_paragraph().paragraph_format.first_line_indent = Cm(0)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --- мутации ---------------------------------------------------------------------------
class Thesis:
    """Обёртка: blocks + мутации. Каждая мутация возвращает НОВЫЙ Thesis."""

    def __init__(self, blocks=None, pages=70, **table_opts):
        self.blocks = blocks if blocks is not None else clean_blocks(pages)
        self.table_opts = table_opts

    def copy(self) -> Thesis:
        return Thesis(copy.deepcopy(self.blocks), **self.table_opts)

    def bytes(self) -> bytes:
        if getattr(self, "_bytes", None) is None:
            self._bytes = render(self.blocks, table_opts=self.table_opts)
        return self._bytes

    # поиск
    def _idx(self, startswith: str, nth: int = 0) -> int:
        want_toc = "........" in startswith
        hits = [
            i
            for i, b in enumerate(self.blocks)
            if (b.get("t") or b.get("caption") or "").startswith(startswith)
            and (want_toc or b.get("tag") != "toc")
        ]
        if len(hits) <= nth:
            raise KeyError(f"нет блока, начинающегося с {startswith!r}")
        return hits[nth]

    def drop(self, startswith: str, nth: int = 0) -> Thesis:
        t = self.copy()
        del t.blocks[t._idx(startswith, nth)]
        return t

    def no_caption(self, startswith: str) -> Thesis:
        """Оставляет таблицу/рисунок, убирая только подпись."""
        t = self.copy()
        blk = t.blocks[t._idx(startswith)]
        blk["caption"] = None
        return t

    def drop_section(self, heading_startswith: str) -> Thesis:
        """Удаляет заголовок и все блоки до следующего h1/h2 включительно не затрагивая его."""
        t = self.copy()
        i = t._idx(heading_startswith)
        j = i + 1
        while j < len(t.blocks) and t.blocks[j]["k"] not in ("h1", "h2"):
            j += 1
        del t.blocks[i:j]
        return t

    def sub(self, old: str, new: str, *, within: str | None = None, count: int = 1) -> Thesis:
        """Заменяет текст; `within` - начало заголовка раздела, внутри которого искать."""
        t = self.copy()
        lo, hi = 0, len(t.blocks)
        if within:
            lo = t._idx(within)
            hi = lo + 1
            while hi < len(t.blocks) and t.blocks[hi]["k"] not in ("h1", "h2"):
                hi += 1
            if t.blocks[lo]["k"] == "h1":  # h1 охватывает и вложенные h2
                hi = lo + 1
                while hi < len(t.blocks) and t.blocks[hi]["k"] != "h1":
                    hi += 1
        done = 0
        for blk in t.blocks[lo:hi]:
            for key in ("t", "caption"):
                if key in blk and old in blk[key] and done < count:
                    blk[key] = blk[key].replace(old, new, 1)
                    done += 1
        if not done:
            raise KeyError(f"{old!r} не найден (within={within})")
        return t

    def insert_after(self, startswith: str, *new: dict) -> Thesis:
        t = self.copy()
        i = t._idx(startswith)
        t.blocks[i + 1 : i + 1] = list(new)
        return t

    def insert_before(self, startswith: str, *new: dict) -> Thesis:
        t = self.copy()
        i = t._idx(startswith)
        t.blocks[i:i] = list(new)
        return t

    def move(self, startswith: str, before: str) -> Thesis:
        t = self.copy()
        blk = t.blocks.pop(t._idx(startswith))
        t.blocks.insert(t._idx(before), blk)
        return t


def mutate(base: Thesis, fn: Callable[[Thesis], Thesis]) -> Thesis:
    return fn(base)


# --- работа для проверки матрицы связности (MAT-01) ---------------------------------------
def matrix_blocks() -> list[dict]:
    """1 гипотеза из 2 частей, 3 положения, 5 задач -> 10 строк. Связи размечены заранее."""
    _COUNTER[0] = 0
    b: list[dict] = [
        p("ВЫПУСКНАЯ КВАЛИФИКАЦИОННАЯ РАБОТА", False, "center"),
        p("Тревожность, самооценка и мотивация студентов", False, "center"),
        h1("СОДЕРЖАНИЕ"),
    ]
    b += [
        p(f"{t} ........ {i + 3}", False, tag="toc")
        for i, t in enumerate(
            (
                "ВВЕДЕНИЕ",
                "ГЛАВА 1. Теоретические основы",
                "1.1. Тревожность студентов",
                "1.2. Самооценка личности",
                "1.3. Учебная мотивация",
                "Выводы по главе 1",
                "ГЛАВА 2. Эмпирическое исследование",
                "2.1. Организация и методы исследования",
                "2.2. Результаты исследования",
                "Выводы по главе 2",
                "ЗАКЛЮЧЕНИЕ",
            )
        )
    ]
    b += [
        h1("ВВЕДЕНИЕ"),
        p("Актуальность исследования определяется ростом интереса к благополучию студентов."),
        p("Цель исследования: изучить тревожность, самооценку и мотивацию студентов."),
        p(
            "Гипотеза исследования: (1) высокая тревожность студентов связана с низкой самооценкой; "
            "(2) мотивация юношей и девушек различается."
        ),
        p("Задачи исследования:"),
        item(1, "Проанализировать литературу по проблеме тревожности, самооценки и мотивации."),
        item(2, "Определить уровень тревожности студентов."),
        item(3, "Определить уровень самооценки студентов."),
        item(4, "Выявить связь тревожности и самооценки студентов."),
        item(5, "Сравнить мотивацию юношей и девушек."),
        p("Положения, выносимые на защиту:"),
        p("1. Тревожность студентов связана с их самооценкой.", False),
        p("2. Мотивация юношей и девушек различается.", False),
        p("3. Уровень тревожности студентов умеренный.", False),
        h1("ГЛАВА 1. Теоретические основы"),
        h2("1.1. Тревожность студентов"),
        p(
            "Тревожность понимается как устойчивая склонность переживать беспокойство. Ч.Д. Спилбергер "
            "выделил ситуативную и личностную тревожность. Тревожность студентов растёт в сессию."
        ),
        filler(30, "тревожности"),
        h2("1.2. Самооценка личности"),
        p(
            "Самооценка личности формируется в юности. С.Р. Пантилеев описал структуру самооценки, "
            "а самооценка студентов связана с тревожностью."
        ),
        filler(30, "самооценки"),
        h2("1.3. Учебная мотивация"),
        p(
            "Мотивация обучения рассматривается как система мотивов. А.А. Реан описал мотивацию студентов, "
            "различия мотивации юношей и девушек обсуждаются в литературе."
        ),
        filler(30, "мотивации"),
        h2("Выводы по главе 1"),
        p(
            "Анализ литературы показал, что тревожность, самооценка и мотивация студентов взаимосвязаны."
        ),
        h1("ГЛАВА 2. Эмпирическое исследование"),
        h2("2.1. Организация и методы исследования"),
        p("Выборку составили 100 человек."),
        p("Тревожность изучалась методикой «Шкала тревожности» (ШТ) Ч.Д. Спилбергера."),
        p("Самооценка изучалась методикой «Опросник самооценки» (ОС) С.Р. Пантилеева."),
        p("Мотивация изучалась методикой «Шкала мотивации обучения» (ШМО) А.А. Реана."),
        p(
            "Данные обрабатывались с помощью коэффициента корреляции Спирмена и U-критерия Манна–Уитни."
        ),
        filler(30, "организации"),
        h2("2.2. Результаты исследования"),
        p("Уровень тревожности студентов представлен в табл. 1."),
        table(
            "Таблица 1 – Уровень тревожности студентов",
            [
                ["Уровень", "Доля"],
                ["Низкий уровень тревожности", "20%"],
                ["Умеренный уровень тревожности", "60%"],
                ["Высокий уровень тревожности", "20%"],
            ],
        ),
        p("Уровень самооценки студентов представлен в табл. 2."),
        table(
            "Таблица 2 – Уровень самооценки студентов",
            [
                ["Уровень", "Доля"],
                ["Низкий уровень самооценки", "25%"],
                ["Средний уровень самооценки", "50%"],
                ["Высокий уровень самооценки", "25%"],
            ],
        ),
        p(
            f"Связь тревожности и самооценки студентов отрицательна (rs{NB}={NB}-0,52; p{NB}<{NB}0,01), табл. 3."
        ),
        table(
            "Таблица 3 – Связь тревожности и самооценки",
            [["Показатель", "rs"], ["Тревожность и самооценка", "-0,52"]],
        ),
        p(f"Мотивация юношей и девушек различается (U{NB}={NB}310,5; p{NB}<{NB}0,05), табл. 4."),
        table(
            "Таблица 4 – Различия мотивации юношей и девушек",
            [["Группа", "Среднее"], ["Юноши", "30,1"], ["Девушки", "35,4"]],
        ),
        p(
            "Гипотеза о связи тревожности и самооценки подтверждена; гипотеза о различиях мотивации подтверждена."
        ),
        filler(30, "результатов"),
        h2("Выводы по главе 2"),
        p(
            "Результаты эмпирического исследования подтверждают гипотезу о связи тревожности и самооценки."
        ),
        h1("ЗАКЛЮЧЕНИЕ"),
        p(
            "1. Анализ литературы по проблеме тревожности, самооценки и мотивации показал их взаимосвязь.",
            False,
        ),
        p("2. Уровень тревожности студентов определён: преобладает умеренный уровень.", False),
        p("3. Уровень самооценки студентов определён: преобладает средний уровень.", False),
        p("4. Связь тревожности и самооценки студентов выявлена: она отрицательна.", False),
        p("5. Мотивация юношей и девушек сравнена: различия значимы.", False),
    ]
    return b
