"""ThesisBuilder: fluent-API для синтетических ВКР поверх python-docx (ARCHITECTURE_TASKS.md §3).

Детерминизм: фиксированные даты в core.xml, фиксированные времена записей zip, текст из
seed. Два вызова с одинаковыми аргументами дают побайтно одинаковый файл.
"""

from __future__ import annotations

import io
import struct
import zipfile
import zlib
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

from .text import TextGen
from .xml_patch import add_complex_field, add_page_break_before

FIXED_DATE = (2026, 1, 1, 0, 0, 0)
BASE_FONT = "Times New Roman"


def _tiny_png() -> bytes:
    """Валидный PNG 8x8 серого цвета без внешних зависимостей."""

    def chunk(kind: bytes, data: bytes) -> bytes:
        body = kind + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body))

    raw = b"".join(b"\x00" + b"\x80" * 8 for _ in range(8))
    ihdr = struct.pack(">IIBBBBB", 8, 8, 8, 0, 0, 0, 0)
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", ihdr)
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


class ThesisBuilder:
    """Собирает .docx: титул, оглавление, введение, главы, заключение, литература, приложения."""

    def __init__(self, seed: int = 1) -> None:
        self.text = TextGen(seed)
        self.doc = Document()
        self._table_no = 0
        self._figure_no = 0
        self._setup_page_and_styles()

    # --- настройка страницы и стилей ---------------------------------------------------
    def _setup_page_and_styles(self) -> None:
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin, sec.right_margin = Cm(3), Cm(1.5)
        sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
        normal = self.doc.styles["Normal"]
        normal.font.name = BASE_FONT
        normal.font.size = Pt(14)
        rpr = normal.element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(_qn(attr), BASE_FONT)
        pf = normal.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        pf.first_line_indent = Cm(1.25)
        pf.space_before = pf.space_after = Pt(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for name, size in (("Heading 1", 14), ("Heading 2", 14)):
            st = self.doc.styles[name]
            st.font.name = BASE_FONT
            st.font.size = Pt(size)
            st.font.bold = True
            st.font.color.rgb = None
            st.element.get_or_add_rPr().get_or_add_rFonts().set(_qn("w:ascii"), BASE_FONT)
            st.element.get_or_add_rPr().get_or_add_rFonts().set(_qn("w:hAnsi"), BASE_FONT)
            st.paragraph_format.first_line_indent = Cm(0)
            st.paragraph_format.keep_with_next = True
            st.paragraph_format.space_before = Pt(0)
            st.paragraph_format.space_after = Pt(0)
        self.doc.styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp = self.doc.core_properties
        cp.author, cp.last_modified_by, cp.title = "Синтетика", "Синтетика", "Тестовая работа"
        cp.created = cp.modified = _fixed_datetime()

    # --- примитивы ---------------------------------------------------------------------
    def para(self, text: str, *, indent: bool = True, align=None, bold: bool = False) -> Paragraph:
        p = self.doc.add_paragraph()
        run = p.add_run(text)
        run.bold = bold or None
        if not indent:
            p.paragraph_format.first_line_indent = Cm(0)
        if align is not None:
            p.alignment = align
        return p

    def heading(self, text: str, level: int = 1, *, new_page: bool = False) -> Paragraph:
        p = self.doc.add_heading(text, level=level)
        if new_page:
            add_page_break_before(p)
        return p

    def body_pages(self, pages: float) -> None:
        for chunk in self.text.paragraphs_for_pages(pages):
            self.para(chunk)

    # --- разделы работы ----------------------------------------------------------------
    def title_page(
        self, title: str = "Влияние учебной мотивации на адаптацию студентов"
    ) -> ThesisBuilder:
        for line in ("МИНИСТЕРСТВО ОБРАЗОВАНИЯ (вымышленное)", "Кафедра психологии"):
            self.para(line, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)
        self.para("ВЫПУСКНАЯ КВАЛИФИКАЦИОННАЯ РАБОТА", indent=False, bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER)  # fmt: skip
        self.para(title, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)
        self.para(f"Выполнил: {self.text.author()}", indent=False, align=WD_ALIGN_PARAGRAPH.RIGHT)
        self.para(f"Научный руководитель: {self.text.author()}", indent=False,
                  align=WD_ALIGN_PARAGRAPH.RIGHT)  # fmt: skip
        return self

    def toc(
        self, entries: list[tuple[str, int]] | None = None, *, field: bool = False
    ) -> ThesisBuilder:
        """Оглавление: текстовое (по умолчанию) или поле TOC с сохранённым результатом."""
        entries = entries or [("ВВЕДЕНИЕ", 3), ("ЗАКЛЮЧЕНИЕ", 60)]
        self.heading("ОГЛАВЛЕНИЕ", 1, new_page=True)
        if field:
            holder = self.para("", indent=False)
            add_complex_field(holder, 'TOC \\o "1-3" \\h', entries[0][0])
            for title, page in entries[1:]:
                self.para(f"{title}\t{page}", indent=False)
        else:
            for title, page in entries:
                self.para(f"{title} ........ {page}", indent=False)
        return self

    def intro(self, elements: dict[str, str] | None = None, pages: float = 3) -> ThesisBuilder:
        self.heading("ВВЕДЕНИЕ", 1, new_page=True)
        elements = elements or {
            "Актуальность": "определяется ростом интереса к вопросам адаптации студентов.",
            "Цель": "выявить связь учебной мотивации и адаптации студентов.",
            "Объект": "учебная мотивация студентов.",
            "Предмет": "связь мотивации и адаптации.",
            "Гипотеза": "учебная мотивация связана с уровнем адаптации студентов.",
        }
        for key, value in elements.items():
            self.para(f"{key} {value}")
        self.body_pages(max(pages - 1, 0.5))
        return self

    def chapter(self, number: int, title: str, subsections: list[str] | None = None,
                pages: float = 20) -> ThesisBuilder:  # fmt: skip
        self.heading(f"ГЛАВА {number}. {title}", 1, new_page=True)
        subsections = subsections or [f"Раздел {number}.1", f"Раздел {number}.2"]
        per = pages / max(len(subsections), 1)
        for i, sub in enumerate(subsections, 1):
            self.heading(f"{number}.{i}. {sub}", 2)
            self.body_pages(per)
        return self

    def table(
        self, rows: list[list[str]], caption: str, *, style: str = "Table Grid"
    ) -> ThesisBuilder:
        self._table_no += 1
        self.para(f"Таблица {self._table_no} – {caption}", indent=False)
        table = self.doc.add_table(rows=len(rows), cols=len(rows[0]), style=style)
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                cell_p = table.cell(r, c).paragraphs[0]
                cell_p.add_run(value)
                cell_p.paragraph_format.first_line_indent = Cm(0)
                cell_p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                cell_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        self.para("", indent=False)
        return self

    def figure(self, caption: str) -> ThesisBuilder:
        self._figure_no += 1
        holder = self.doc.add_paragraph()
        holder.alignment = WD_ALIGN_PARAGRAPH.CENTER
        holder.paragraph_format.first_line_indent = Cm(0)
        holder.add_run().add_picture(io.BytesIO(_tiny_png()), width=Cm(4))
        self.para(f"Рисунок {self._figure_no} – {caption}", indent=False,
                  align=WD_ALIGN_PARAGRAPH.CENTER)  # fmt: skip
        return self

    def conclusion(self, pages: float = 2) -> ThesisBuilder:
        self.heading("ЗАКЛЮЧЕНИЕ", 1, new_page=True)
        self.body_pages(pages)
        return self

    def bibliography(self, n: int = 30) -> ThesisBuilder:
        self.heading("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", 1, new_page=True)
        for i in range(1, n + 1):
            self.para(f"{i}. {self.text.bib_entry(i)}", indent=False)
        return self

    def appendix(self, key: str, title: str = "Вспомогательные материалы") -> ThesisBuilder:
        self.heading(f"ПРИЛОЖЕНИЕ {key}", 1, new_page=True)
        self.para(title, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)
        self.body_pages(0.5)
        return self

    # --- сохранение --------------------------------------------------------------------
    def to_bytes(self) -> bytes:
        buf = io.BytesIO()
        self.doc.save(buf)
        return _normalize_zip(buf.getvalue())

    def save(self, path: Path) -> Path:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(self.to_bytes())
        return path


def _qn(tag: str) -> str:
    from docx.oxml.ns import qn

    return qn(tag)


def _fixed_datetime():
    from datetime import datetime

    return datetime(*FIXED_DATE)


def _normalize_zip(data: bytes) -> bytes:
    """Фиксированные времена записей и порядок: побайтовая воспроизводимость."""
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as zin,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout,
    ):
        for info in sorted(zin.infolist(), key=lambda i: i.filename):
            zi = zipfile.ZipInfo(info.filename, date_time=FIXED_DATE)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            zout.writestr(zi, zin.read(info.filename))
    return out.getvalue()
