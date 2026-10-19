import hashlib
import re

import pytest
from docx import Document

from tests.fixtures.builders import TextGen, ThesisBuilder, variants
from tests.fixtures.builders.base import _tiny_png
from tests.fixtures.builders.registry import BUILDERS, HEAVY_FIXTURES
from tests.fixtures.builders.text import CHARS_PER_PAGE
from vkr_checker.docx.package import DocxError, DocxPackage


def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def test_text_is_deterministic_per_seed():
    a, b, c = TextGen(1), TextGen(1), TextGen(2)
    assert a.paragraph() == b.paragraph()
    assert TextGen(1).paragraph() != c.paragraph()


def test_text_volume_close_to_requested_pages():
    chars = sum(len(p) for p in TextGen(3).paragraphs_for_pages(10))
    assert 10 * CHARS_PER_PAGE <= chars <= 10 * CHARS_PER_PAGE + 1000


def test_authors_are_from_fictional_list():
    gen = TextGen(5)
    for _ in range(20):
        assert re.fullmatch(r"[А-ЯЁ][а-яё]+ [А-Я]\.[А-Я]\.", gen.author())


def test_builder_output_is_byte_deterministic():
    assert sha(variants.clean(8).to_bytes()) == sha(variants.clean(8).to_bytes())
    assert sha(variants.clean(8, seed=1).to_bytes()) != sha(variants.clean(8, seed=2).to_bytes())


def test_png_is_valid_png():
    assert _tiny_png()[:8] == b"\x89PNG\r\n\x1a\n"


def test_clean_has_all_sections_tables_figures(tmp_path):
    path = variants.clean(10).save(tmp_path / "c.docx")
    doc = Document(str(path))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    for needle in (
        "ОГЛАВЛЕНИЕ",
        "ВВЕДЕНИЕ",
        "ГЛАВА 1",
        "ГЛАВА 2",
        "ЗАКЛЮЧЕНИЕ",
        "СПИСОК",
        "ПРИЛОЖЕНИЕ 1",
    ):
        assert any(h.startswith(needle) for h in headings), needle
    assert len(doc.tables) == 2
    captions = [p.text for p in doc.paragraphs if re.match(r"(Таблица|Рисунок) \d+ – ", p.text)]
    assert captions == ["Таблица 1 – Описательные статистики", "Рисунок 1 – Схема исследования",
                        "Таблица 2 – Сравнение групп"]  # fmt: skip
    assert len(doc.inline_shapes) == 1


def test_clean_page_geometry_and_font(tmp_path):
    doc = Document(str(variants.clean(5).save(tmp_path / "c.docx")))
    sec = doc.sections[0]
    assert round(sec.left_margin.cm, 1) == 3.0 and round(sec.right_margin.cm, 1) == 1.5
    assert doc.styles["Normal"].font.name == "Times New Roman"
    assert doc.styles["Normal"].font.size.pt == 14


@pytest.mark.parametrize("kind", ["INTRO", "CONCLUSION", "BIBLIOGRAPHY", "APPENDIX"])
def test_missing_section_removes_exactly_that_heading(tmp_path, kind):
    doc = Document(str(variants.missing_section(kind, 6).save(tmp_path / "m.docx")))
    marker = {"INTRO": "ВВЕДЕНИЕ", "CONCLUSION": "ЗАКЛЮЧЕНИЕ", "BIBLIOGRAPHY": "СПИСОК",
              "APPENDIX": "ПРИЛОЖЕНИЕ"}[kind]  # fmt: skip
    heads = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert not any(h.startswith(marker) for h in heads)
    assert len(heads) == 6  # остальные разделы на месте


def test_table_indent_mutation_sets_direct_indent(tmp_path):
    doc = Document(str(variants.table_indent(6, cm=1.25).save(tmp_path / "t.docx")))
    p = doc.tables[0].cell(0, 0).paragraphs[0]
    assert round(p.paragraph_format.first_line_indent.cm, 2) == 1.25


def test_toc_field_keeps_stale_result(tmp_path):
    b = ThesisBuilder().title_page().toc([("ВВЕДЕНИЕ", 3), ("ЗАКЛЮЧЕНИЕ", 9)], field=True)
    xml = b.to_bytes()
    with DocxPackage.open(b.save(tmp_path / "f.docx")) as pkg:
        body = pkg.read("word/document.xml").decode()
    assert xml
    assert "TOC" in body and 'fldCharType="separate"' in body


def test_generated_good_docx_accepted_and_bad_ones_rejected(tmp_path):
    for name in ("clean_20p", "rule_STR-01_no_conclusion"):
        p = tmp_path / f"{name}.docx"
        p.write_bytes(BUILDERS[name]())
        DocxPackage.open(p).close()
    expected = {
        "bad_pdf_as_docx": "unsupported_format",
        "bad_doc_as_docx": "unsupported_format",
        "bad_empty": "unsupported_format",
        "bad_odt_as_docx": "unsupported_format",
    }
    for name, code in expected.items():
        p = tmp_path / f"{name}.docx"
        p.write_bytes(BUILDERS[name]())
        with pytest.raises(DocxError) as exc:
            DocxPackage.open(p)
        assert exc.value.code == code, name
    assert "bad_zipbomb" in HEAVY_FIXTURES
