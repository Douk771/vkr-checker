from io import BytesIO

from openpyxl import load_workbook

from vkr_checker.matrix.xlsx import AUTO_PREFIX, build_matrix_xlsx, export_filename
from vkr_checker.model.finding import Location
from vkr_checker.report.schema import Matrix, MatrixCells, MatrixRow, MatrixStats


def cell(status, text="", cands=(), locs=()):
    return dict(status=status, text=text, candidates=list(cands), locations=list(locs))


def make_matrix():
    loc = Location(section_path=["Глава 1", "1.1 Теория"], paragraph_idx=7, quote="цитата")
    cells = MatrixCells(
        theory=cell("auto", "1.1, Иванов", locs=[loc]),
        instrument=cell("gap", cands=["MBI"]),
        calculation=cell("n/a"),
        result=cell("gap"),
        conclusion=cell("auto", "вывод"),
    )
    row = MatrixRow(kind="task", label="Задача 1", statement="Изучить X",
                    statement_loc=loc, is_theoretical=False, cells=cells)  # fmt: skip
    return Matrix(rows=[row], stats=MatrixStats(rows=1, gaps=2, auto_cells=2, warnings=[]))


def test_workbook_content():
    data = build_matrix_xlsx(make_matrix(), "работа.docx", "1.0")
    wb = load_workbook(BytesIO(data))
    ws = wb["Матрица"]
    assert ws.freeze_panes == "A2"
    row = [c.value for c in ws[2]]
    assert row[0] == "Задача 1"
    assert row[2].startswith(AUTO_PREFIX)
    assert row[3].startswith("РАЗРЫВ") and "MBI" in row[3]
    assert row[4] == "не требуется"
    assert row[7] == "Задача"
    assert row[8] == "есть разрывы: 2"
    assert ws.cell(2, 3).font.italic
    places = wb["Места"]
    assert places.max_row == 3  # заголовок + формулировка + теория
    assert places.cell(2, 6).value == "цитата"


def test_control_chars_are_stripped():
    m = make_matrix()
    m.rows[0].statement = "a\x00b\x0bc"
    wb = load_workbook(BytesIO(build_matrix_xlsx(m, "f.docx")))
    assert wb["Матрица"].cell(2, 2).value == "abc"


def test_export_filename_sanitized():
    assert export_filename("../../etc/passwd.docx") == "etc_passwd_матрица.xlsx"
    assert export_filename("Иванов ВКР.docx") == "Иванов ВКР_матрица.xlsx"
    assert export_filename("...docx") == "vkr_матрица.xlsx"


def test_leading_equals_is_text_not_formula():
    m = make_matrix()
    m.rows[0].statement = "=1+1"
    m.rows[0].label = "@x"
    wb = load_workbook(BytesIO(build_matrix_xlsx(m, "f.docx")))
    cell = wb["Матрица"].cell(2, 2)
    assert cell.value == "=1+1" and cell.data_type == "s"
