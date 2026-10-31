import json
from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

from vkr_checker.config import Settings
from vkr_checker.main import create_app
from vkr_checker.report.docx_export import build_report_docx, report_filename
from vkr_checker.report.schema import CheckReport, CompareReport

GOLDEN = Path(__file__).resolve().parents[1] / "golden"


def golden(name: str) -> dict:
    return json.loads((GOLDEN / name).read_text(encoding="utf-8"))


def all_text(data: bytes) -> str:
    doc = Document(BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        parts += [c.text for row in t.rows for c in row.cells]
    return "\n".join(parts)


def test_check_report_docx_contains_summary_and_findings():
    report = CheckReport.model_validate(golden("check_findings.json"))
    text = all_text(build_report_docx(check=report))
    assert "Отчёт о проверке ВКР" in text
    assert report.file.name in text
    findings = [
        f for b in report.blocks for r in b.rule_results if r.rule_id != "NC-01"
        for f in r.findings if f.level != "not_checked"
    ]  # fmt: skip
    assert findings
    for f in findings:
        assert f.message in text
    assert "Нарушение." in text or "Подозрение." in text
    for item in report.not_checked:
        assert item.title in text


def test_clean_report_says_nothing_found():
    report = CheckReport.model_validate(golden("check_clean.json"))
    assert "Нарушений и подозрений не найдено." in all_text(build_report_docx(check=report))


def test_compare_report_docx_has_diff_and_new_version_check():
    cmp = CompareReport.model_validate(golden("compare.json"))
    text = all_text(build_report_docx(compare=cmp))
    assert "Сравнение версий ВКР" in text
    assert cmp.old.file.name in text and cmp.new.file.name in text
    assert "Динамика замечаний" in text
    assert f"Проверка новой версии: {cmp.new.file.name}" in text


def test_report_filename():
    assert report_filename("Моя работа v2.docx", compare=False) == "Моя работа v2_отчёт.docx"
    assert report_filename("a/b.docx", compare=True) == "a_b_сравнение.docx"


def client() -> TestClient:
    return TestClient(create_app(Settings(mode="local")))


def test_export_endpoint_returns_docx_for_check_and_compare():
    c = client()
    r = c.post("/api/export/report.docx", json={"check": golden("check_findings.json")})
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert "_%D0%BE%D1%82%D1%87%D1%91%D1%82.docx" in r.headers["content-disposition"]  # _отчёт
    assert "Отчёт о проверке ВКР" in all_text(r.content)
    r = c.post("/api/export/report.docx", json={"compare": golden("compare.json")})
    assert r.status_code == 200 and "Сравнение версий ВКР" in all_text(r.content)


def test_export_endpoint_requires_exactly_one_report():
    c = client()
    for body in ({}, {"check": golden("check_clean.json"), "compare": golden("compare.json")}):
        r = c.post("/api/export/report.docx", json=body)
        assert r.status_code == 422 and r.json()["error"]["code"] == "validation_error"


def test_rewritten_phrase_grouped_into_one_deletion_and_one_insertion():
    from vkr_checker.report.docx_export import _group_segments
    from vkr_checker.report.schema import Segment

    segs = [Segment(op=o, text=t) for o, t in [
        ("del", "Учебная"), ("ins", "Выявлена"), ("equal", " "),
        ("del", "мотивация"), ("ins", "связь"), ("equal", " (r = 0,45)"),
    ]]  # fmt: skip
    got = [(s.op, s.text) for s in _group_segments(segs)]
    assert got == [
        ("del", "Учебная мотивация"),
        ("ins", "Выявлена связь"),
        ("equal", " (r = 0,45)"),
    ]
