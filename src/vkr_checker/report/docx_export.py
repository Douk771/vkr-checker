"""Выгрузка отчёта в .docx: проверка одной работы или сравнение двух версий.

Отчёт строится из присланного клиентом JSON (сервер ничего не хранит), как и .xlsx матрицы.
Уровень находки подписан словом, поэтому документ читается и без цвета, и при печати.
"""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from vkr_checker.model.finding import Finding, Location
from vkr_checker.report.schema import CheckReport, CompareReport, Segment

LEVEL_TITLES = {
    "violation": "Нарушение",
    "suspicion": "Подозрение",
    "not_checked": "Не проверяется",
}
LEVEL_COLORS = {
    "violation": RGBColor(0xB0, 0x1E, 0x1E),
    "suspicion": RGBColor(0x9A, 0x6A, 0x00),
    "not_checked": RGBColor(0x55, 0x55, 0x55),
}
TREND_TITLES = {"less": "стало меньше", "more": "стало больше", "same": "без изменений"}
DEL_COLOR = RGBColor(0xB0, 0x1E, 0x1E)
INS_COLOR = RGBColor(0x1E, 0x6B, 0x2E)
MAX_LOCATIONS = 5  # остальные места — в отчёте .html и в интерфейсе
# Word не принимает управляющие символы в тексте
_ILLEGAL = {c: None for c in range(32) if c not in (9, 10, 13)}


def _t(text: object) -> str:
    return str(text if text is not None else "").translate(_ILLEGAL)


def _new_document():
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin, sec.right_margin = Cm(2.5), Cm(1.5)
    sec.top_margin = sec.bottom_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(4)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        _set_font(style, "Times New Roman")
        style.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)
    return doc


def _set_font(style, name: str) -> None:
    """Шрифт стиля без шрифта темы: заголовки шаблона python-docx ссылаются на тему."""
    rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)


def _meta(doc, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = _t(key), _t(value)
    _fix_widths(table, (Cm(8.5), Cm(8.5)))
    doc.add_paragraph()


def _fix_widths(table, widths) -> None:
    """Фиксированные ширины столбцов: и в сетке таблицы, и в каждой ячейке (иначе Word и
    LibreOffice раскладывают столбцы поровну)."""
    table.autofit = False
    for col, width in zip(table.columns, widths, strict=True):
        col.width = width
    for row in table.rows:
        for cell, width in zip(row.cells, widths, strict=True):
            cell.width = width


def _location_text(loc: Location) -> str:
    parts = []
    if loc.section_path:
        parts.append(" › ".join(loc.section_path))
    if loc.table is not None:
        cell = loc.table.number_label
        if loc.table.row is not None and loc.table.col is not None:
            cell += f", строка {loc.table.row}, столбец {loc.table.col}"
        parts.append(cell)
    elif loc.paragraph_idx is not None:
        parts.append(f"абзац {loc.paragraph_idx}")
    if loc.page is not None:
        parts.append(f"≈ стр. {loc.page} (вёрстка LibreOffice)")
    return "; ".join(parts)


def _finding(doc, f: Finding) -> None:
    p = doc.add_paragraph(style="List Bullet")
    label = p.add_run(f"{LEVEL_TITLES.get(f.level, f.level)}. ")
    label.bold = True
    label.font.color.rgb = LEVEL_COLORS.get(f.level, RGBColor(0, 0, 0))
    p.add_run(_t(f.message))
    for loc in f.locations[:MAX_LOCATIONS]:
        where = _location_text(loc)
        lp = doc.add_paragraph()
        lp.paragraph_format.left_indent = Cm(1)
        if where:
            run = lp.add_run(where)
            run.italic = True
        if loc.quote:
            lp.add_run(f"{chr(10) if where else ''}«{_t(loc.quote)}»")
    if len(f.locations) > MAX_LOCATIONS:
        more = doc.add_paragraph(f"…и ещё мест: {len(f.locations) - MAX_LOCATIONS}")
        more.paragraph_format.left_indent = Cm(1)
    if f.hint:
        hp = doc.add_paragraph()
        hp.paragraph_format.left_indent = Cm(1)
        run = hp.add_run(f"Подсказка: {_t(f.hint)}")
        run.font.size = Pt(11)


def _reportable(findings: list[Finding]) -> bool:
    """Есть что показать: нарушение, подозрение или проверка, которая не выполнилась."""
    return any(f.level != "not_checked" or f.not_run for f in findings)


def _summary_rows(report: CheckReport) -> list[tuple[str, str]]:
    s = report.summary
    pages = report.volume.pages_main
    volume = "—" if pages is None else f"{pages:g} стр."
    volume += " (по вёрстке LibreOffice)" if report.volume.method == "layout" else " (оценка)"
    return [
        ("Файл", report.file.name),
        ("Проверено", report.checked_at.strftime("%d.%m.%Y %H:%M") + " UTC"),
        ("Версия сервиса", report.service_version),
        ("Основная часть", volume),
        ("Нарушений", str(s.violation)),
        ("Подозрений", str(s.suspicion)),
        ("Не проверяется кодом", str(s.not_checked)),
        ("Пройдено правил", str(len(s.passed_rules))),
    ]


def _check_body(doc, report: CheckReport, heading_level: int) -> None:
    _meta(doc, _summary_rows(report))
    notes = [*report.warnings]
    if report.layout.disclaimer and report.volume.method == "layout":
        notes.append(report.layout.disclaimer)
    if report.summary.noise_warning:
        notes.append(report.summary.noise_warning)
    for note in notes:
        doc.add_paragraph(_t(note)).runs[0].italic = True

    any_findings = False
    for block in report.blocks:
        rules = [r for r in block.rule_results if r.rule_id != "NC-01" and _reportable(r.findings)]
        if not rules:
            continue
        any_findings = True
        doc.add_heading(_t(block.title), level=heading_level)
        for rule in rules:
            doc.add_heading(f"{rule.rule_id} {_t(rule.title)}", level=heading_level + 1)
            for f in rule.findings:
                _finding(doc, f)
    if not any_findings:
        doc.add_paragraph("Нарушений и подозрений не найдено.")

    if report.not_checked:
        doc.add_heading("Не проверяется автоматически", level=heading_level)
        for item in report.not_checked:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(f"{_t(item.title)}").bold = True
            p.add_run(f" — {_t(item.source)}")
            if item.hint:
                p.add_run(f". {_t(item.hint)}")
    if report.matrix.rows:
        doc.add_paragraph(
            "Черновик матрицы связности выгружается отдельно в .xlsx (кнопка «Скачать .xlsx»)."
        ).runs[0].italic = True


def _group_segments(segments: list[Segment]) -> list[Segment]:
    """Пословный дифф переписанной фразы склеивается в один блок «удалено» и один «добавлено»
    (как во фронтенде): пробел между изменениями входит в оба блока, только когда в текущем
    блоке уже есть и удаление, и вставка."""
    out: list[Segment] = []
    deleted = inserted = ""

    def flush() -> None:
        nonlocal deleted, inserted
        if deleted:
            out.append(Segment(op="del", text=deleted))
        if inserted:
            out.append(Segment(op="ins", text=inserted))
        deleted = inserted = ""

    for i, seg in enumerate(segments):
        nxt = segments[i + 1] if i + 1 < len(segments) else None
        if seg.op == "del":
            deleted += seg.text
        elif seg.op == "ins":
            inserted += seg.text
        elif (
            deleted and inserted and not seg.text.strip() and nxt is not None and nxt.op != "equal"
        ):
            deleted += seg.text
            inserted += seg.text
        else:
            flush()
            out.append(seg)
    flush()
    return out


def _segments(paragraph, segments: list[Segment]) -> None:
    for s in _group_segments(segments):
        run = paragraph.add_run(_t(s.text))
        if s.op == "del":
            run.font.strike = True
            run.font.color.rgb = DEL_COLOR
        elif s.op == "ins":
            run.font.underline = True
            run.font.color.rgb = INS_COLOR


def _compare_body(doc, cmp: CompareReport) -> None:
    diff = cmp.diff
    changed = [s for s in diff.sections if s.status != "unchanged"]
    old_s = cmp.old.summary
    _meta(doc, [
        ("Старая версия", cmp.old.file.name),
        ("Новая версия", cmp.new.file.name),
        ("Сравнено", cmp.checked_at.strftime("%d.%m.%Y %H:%M") + " UTC"),
        ("Старая версия: нарушений / подозрений", f"{old_s.violation} / {old_s.suspicion}"),
        ("Новая версия: нарушений / подозрений",
         f"{cmp.new.summary.violation} / {cmp.new.summary.suspicion}"),
        ("Изменено разделов", str(len(changed))),
        ("Изменений чисел", str(len(diff.number_changes))),
    ])  # fmt: skip

    doc.add_heading("Изменения по разделам", level=1)
    if not changed:
        doc.add_paragraph("Различий в тексте не найдено.")
    for sec in changed:
        title = sec.new_title or sec.old_title or "(без заголовка)"
        if sec.status == "renamed" and sec.old_title and sec.new_title:
            title = f"{sec.old_title} → {sec.new_title}"
        doc.add_heading(_t(title), level=2)
        for para in sec.paragraphs:
            if para.status == "unchanged":
                continue
            p = doc.add_paragraph()
            label = {
                "modified": "Изменён", "inserted": "Добавлен", "deleted": "Удалён",
                "moved_out": "Перенесён отсюда", "moved_in": "Перенесён сюда",
            }.get(para.status, para.status)  # fmt: skip
            p.add_run(f"{label}: ").bold = True
            if para.segments:
                _segments(p, para.segments)
            elif para.text:
                run = p.add_run(_t(para.text))
                run.font.strike = para.status == "deleted"

    if diff.number_changes:
        doc.add_heading("Изменения чисел", level=1)
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        for cell, text in zip(table.rows[0].cells, ("Было", "Стало", "Контекст"), strict=True):
            cell.text = text
            cell.paragraphs[0].runs[0].bold = True
        for n in diff.number_changes:
            cells = table.add_row().cells
            cells[0].text, cells[1].text, cells[2].text = _t(n.old), _t(n.new), _t(n.context)
        _fix_widths(table, (Cm(2.5), Cm(2.5), Cm(12)))
        doc.add_paragraph()

    if diff.impact:
        doc.add_heading("Влияние изменений на выводы", level=1)
        for f in diff.impact:
            _finding(doc, f)

    d = diff.dynamics
    doc.add_heading("Динамика замечаний", level=1)
    for title, items in (("Исправлено", d.fixed), ("Появилось", d.new), ("Осталось", d.remaining)):
        doc.add_paragraph(f"{title}: {len(items)}").runs[0].bold = True
        for item in items:
            trend = f" ({TREND_TITLES[item.trend]})" if item.trend else ""
            doc.add_paragraph(
                f"{item.rule_id} {_t(item.title)}{trend}: {_t(item.message)}", style="List Bullet"
            )

    doc.add_page_break()
    doc.add_heading(f"Проверка новой версии: {_t(cmp.new.file.name)}", level=1)
    _check_body(doc, cmp.new, heading_level=2)


def build_report_docx(
    check: CheckReport | None = None, compare: CompareReport | None = None
) -> bytes:
    """Собирает .docx отчёта о проверке или о сравнении версий."""
    doc = _new_document()
    if compare is not None:
        title = doc.add_heading("Сравнение версий ВКР", level=0)
        _set_center(title)
        _compare_body(doc, compare)
    elif check is not None:
        title = doc.add_heading("Отчёт о проверке ВКР", level=0)
        _set_center(title)
        _check_body(doc, check, heading_level=1)
    else:
        raise ValueError("нужен check или compare")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _set_center(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def report_filename(file_name: str, compare: bool) -> str:
    """Безопасное имя скачиваемого файла: <имя без расширения>_отчёт.docx / _сравнение.docx."""
    stem = file_name.rsplit(".", 1)[0] if "." in file_name else file_name
    safe = "".join(ch if ch.isalnum() or ch in "-_ " else "_" for ch in stem).strip(" _")
    return f"{safe[:80] or 'vkr'}_{'сравнение' if compare else 'отчёт'}.docx"
