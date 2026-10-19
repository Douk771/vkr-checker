"""Выгрузка матрицы связности в .xlsx (ARCHITECTURE_CHECKS.md §7, «Скачивание»).

Статус ячейки подписан словом, поэтому файл читается и без цвета (SPEC §5).
"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from vkr_checker.report.schema import Matrix, MatrixCell, MatrixRow

AUTO_PREFIX = "[авто, проверить] "
GAP_TEXT = "РАЗРЫВ"
NA_TEXT = "не требуется"
KIND_TITLES = {
    "hypothesis_part": "Часть гипотезы",
    "defense_position": "Положение на защиту",
    "task": "Задача",
}
CELL_COLUMNS = (
    ("theory", "Теория"),
    ("instrument", "Методика"),
    ("calculation", "Расчёт"),
    ("result", "Результат"),
    ("conclusion", "Вывод"),
)
AUTO_FILL = PatternFill("solid", fgColor="FFF2CC")
GAP_FILL = PatternFill("solid", fgColor="F4B6B6")
HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
WRAP = Alignment(wrap_text=True, vertical="top")
SHEET_MATRIX = "Матрица"
SHEET_PLACES = "Места"
# Excel не принимает управляющие символы в ячейках
_ILLEGAL = {c: None for c in range(32) if c not in (9, 10, 13)}


def clean(text: str) -> str:
    return text.translate(_ILLEGAL)


def _cell_text(cell: MatrixCell) -> str:
    if cell.status == "gap":
        suffix = f" (кандидаты: {'; '.join(cell.candidates)})" if cell.candidates else ""
        return GAP_TEXT + suffix
    if cell.status == "n/a":
        return NA_TEXT
    return AUTO_PREFIX + cell.text


def _style_cell(target, cell: MatrixCell) -> None:
    target.alignment = WRAP
    if cell.status == "auto":
        target.font = Font(italic=True)
        target.fill = AUTO_FILL
    elif cell.status == "gap":
        target.fill = GAP_FILL
        target.font = Font(bold=True)
    else:
        target.font = Font(color="808080")


def row_status(row: MatrixRow) -> str:
    cells = [getattr(row.cells, key) for key, _ in CELL_COLUMNS]
    gaps = sum(c.status == "gap" for c in cells)
    return f"есть разрывы: {gaps}" if gaps else "разрывов нет"


def _write_matrix(ws: Worksheet, matrix: Matrix) -> None:
    ws.title = SHEET_MATRIX
    header = ["Элемент", "Формулировка", *[t for _, t in CELL_COLUMNS], "Тип", "Статус строки"]
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HEADER_FILL
        c.alignment = WRAP
    for row in matrix.rows:
        ws.append(
            [
                clean(row.label),
                clean(row.statement),
                *[clean(_cell_text(getattr(row.cells, key))) for key, _ in CELL_COLUMNS],
                KIND_TITLES.get(row.kind, row.kind),
                row_status(row),
            ]
        )
        r = ws.max_row
        for col in range(1, len(header) + 1):
            ws.cell(r, col).alignment = WRAP
        for offset, (key, _) in enumerate(CELL_COLUMNS):
            _style_cell(ws.cell(r, 3 + offset), getattr(row.cells, key))
    widths = [18, 45, 32, 32, 32, 32, 32, 20, 20]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def _write_places(ws: Worksheet, matrix: Matrix) -> None:
    ws.title = SHEET_PLACES
    ws.append(["Элемент", "Столбец", "Путь по заголовкам", "Абзац", "Таблица", "Цитата"])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HEADER_FILL
    titles = dict(CELL_COLUMNS)
    for row in matrix.rows:
        entries = []
        if row.statement_loc is not None:
            entries.append(("Формулировка", row.statement_loc))
        for key, title in titles.items():
            entries.extend((title, loc) for loc in getattr(row.cells, key).locations)
        for column, loc in entries:
            ws.append(
                [
                    clean(row.label),
                    column,
                    clean(" > ".join(loc.section_path)),
                    loc.paragraph_idx,
                    clean(loc.table.number_label) if loc.table else None,
                    clean(loc.quote),
                ]
            )
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.alignment = WRAP
    for i, w in enumerate([18, 14, 40, 10, 22, 60], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def build_matrix_xlsx(
    matrix: Matrix,
    file_name: str,
    service_version: str | None = None,
    checked_at: datetime | None = None,
) -> bytes:
    wb = Workbook()
    _write_matrix(wb.active, matrix)
    _write_places(wb.create_sheet(), matrix)
    props = wb.properties
    props.title = clean(f"Матрица связности: {file_name}")[:250]
    props.creator = f"vkr-checker {service_version}" if service_version else "vkr-checker"
    if checked_at is not None:
        props.created = checked_at.replace(tzinfo=None)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_filename(file_name: str) -> str:
    """Безопасное имя скачиваемого файла: <имя без расширения>_matrix.xlsx."""
    stem = file_name.rsplit(".", 1)[0] if "." in file_name else file_name
    safe = "".join(ch if ch.isalnum() or ch in "-_ " else "_" for ch in stem).strip(" _")
    return f"{safe[:80] or 'vkr'}_matrix.xlsx"
